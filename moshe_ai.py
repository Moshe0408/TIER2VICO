"""
MosheAI - סוכן AI חכם לדוחות, שקופיות וסטטיסטיקה
מתחבר ל-Claude API ולומד מטעויות על הדרך
"""

import os
import json
import datetime
import traceback
import anthropic

from pathlib import Path

# ───── תלויות אופציונליות (נטען רק אם קיים) ─────
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Pt as DocxPt, RGBColor as DocxRGB, Inches as DocxInches
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.ticker as mticker
    MATPLOTLIB_AVAILABLE = True
except ImportError:
    MATPLOTLIB_AVAILABLE = False

# ─────────────────────────────────────────────────
MEMORY_FILE = Path(__file__).parent / "moshe_ai_memory.json"
OUTPUT_DIR  = Path(__file__).parent / "moshe_ai_outputs"
OUTPUT_DIR.mkdir(exist_ok=True)

MODEL = "claude-opus-4-6"

SYSTEM_PROMPT = """אתה MosheAI - עוזר AI מקצועי בעברית.
תפקידך:
1. ליצור מצגות PowerPoint מקצועיות
2. ליצור דוחות Word
3. לנתח נתונים ולהציג סטטיסטיקה עם גרפים
4. ללמוד מטעויות קודמות ולהשתפר

כאשר נתבקש, השתמש בכלים הזמינים לך.
ענה תמיד בעברית אלא אם התבקשת אחרת.
"""

# ═══════════════════════════════════════════════
#  מערכת זיכרון ולמידה
# ═══════════════════════════════════════════════

def load_memory() -> dict:
    if MEMORY_FILE.exists():
        try:
            return json.loads(MEMORY_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {"sessions": [], "errors": [], "improvements": [], "stats": {"total_tasks": 0, "success": 0, "failed": 0}}


def save_memory(memory: dict):
    MEMORY_FILE.write_text(json.dumps(memory, ensure_ascii=False, indent=2), encoding="utf-8")


def log_success(memory: dict, task: str, output_path: str):
    memory["stats"]["total_tasks"] += 1
    memory["stats"]["success"] += 1
    memory["sessions"].append({
        "ts": datetime.datetime.now().isoformat(),
        "task": task[:200],
        "status": "success",
        "output": output_path
    })
    # שמור רק 50 session אחרונים
    memory["sessions"] = memory["sessions"][-50:]
    save_memory(memory)


def log_error(memory: dict, task: str, error: str, context: str = ""):
    memory["stats"]["total_tasks"] += 1
    memory["stats"]["failed"] += 1
    entry = {
        "ts": datetime.datetime.now().isoformat(),
        "task": task[:200],
        "error": error[:500],
        "context": context[:300]
    }
    memory["errors"].append(entry)
    memory["errors"] = memory["errors"][-30:]
    save_memory(memory)


def build_lessons_prompt(memory: dict) -> str:
    """בונה תזכורת לקלוד מהטעויות הקודמות"""
    if not memory["errors"]:
        return ""
    recent = memory["errors"][-5:]
    lines = ["לקחים מטעויות קודמות (למד מהן והימנע):"]
    for e in recent:
        lines.append(f"- משימה: {e['task'][:80]} | שגיאה: {e['error'][:120]}")
    return "\n".join(lines)


# ═══════════════════════════════════════════════
#  כלים (Tools)
# ═══════════════════════════════════════════════

TOOLS = [
    {
        "name": "create_presentation",
        "description": "יוצר מצגת PowerPoint (.pptx) עם שקופיות. מקבל כותרת ורשימת שקופיות.",
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "כותרת המצגת"},
                "filename": {"type": "string", "description": "שם הקובץ (ללא סיומת)"},
                "slides": {
                    "type": "array",
                    "description": "רשימת שקופיות",
                    "items": {
                        "type": "object",
                        "properties": {
                            "heading": {"type": "string", "description": "כותרת השקופית"},
                            "bullets": {"type": "array", "items": {"type": "string"}, "description": "נקודות תוכן"},
                            "notes": {"type": "string", "description": "הערות מרצה (אופציונלי)"}
                        },
                        "required": ["heading", "bullets"]
                    }
                },
                "theme_color": {"type": "string", "description": "צבע ראשי בפורמט hex, למשל #1F4E79"}
            },
            "required": ["title", "filename", "slides"]
        }
    },
    {
        "name": "create_word_report",
        "description": "יוצר דוח Word (.docx) עם כותרות, פסקאות וטבלאות.",
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string"},
                "filename": {"type": "string"},
                "sections": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "heading": {"type": "string"},
                            "content": {"type": "string"},
                            "table": {
                                "type": "object",
                                "description": "טבלה אופציונלית",
                                "properties": {
                                    "headers": {"type": "array", "items": {"type": "string"}},
                                    "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}}
                                }
                            }
                        },
                        "required": ["heading", "content"]
                    }
                }
            },
            "required": ["title", "filename", "sections"]
        }
    },
    {
        "name": "create_statistics_chart",
        "description": "יוצר גרף סטטיסטי (עמודות/עוגה/קו) ושומר כתמונה PNG.",
        "input_schema": {
            "type": "object",
            "properties": {
                "chart_type": {"type": "string", "enum": ["bar", "pie", "line", "horizontal_bar"], "description": "סוג הגרף"},
                "title": {"type": "string"},
                "filename": {"type": "string"},
                "labels": {"type": "array", "items": {"type": "string"}},
                "values": {"type": "array", "items": {"type": "number"}},
                "xlabel": {"type": "string"},
                "ylabel": {"type": "string"},
                "colors": {"type": "array", "items": {"type": "string"}, "description": "רשימת צבעים hex (אופציונלי)"}
            },
            "required": ["chart_type", "title", "filename", "labels", "values"]
        }
    },
    {
        "name": "recall_memory",
        "description": "מחזיר סיכום של היסטוריית העבודה, טעויות קודמות ולקחים שנלמדו.",
        "input_schema": {
            "type": "object",
            "properties": {},
            "required": []
        }
    }
]


# ═══════════════════════════════════════════════
#  מימוש הכלים
# ═══════════════════════════════════════════════

def tool_create_presentation(args: dict) -> dict:
    if not PPTX_AVAILABLE:
        return {"error": "python-pptx לא מותקן. הרץ: pip install python-pptx"}

    title    = args["title"]
    filename = args["filename"].rstrip(".pptx") + ".pptx"
    slides   = args["slides"]
    hex_color = args.get("theme_color", "#1F4E79").lstrip("#")

    try:
        r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    except Exception:
        r, g, b = 31, 78, 121

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # ── שקופית כותרת ──
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(r, g, b)
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    if slide.placeholders[1]:
        slide.placeholders[1].text = datetime.datetime.now().strftime("%d/%m/%Y")

    # ── שקופיות תוכן ──
    content_layout = prs.slide_layouts[1]
    for s in slides:
        sl = prs.slides.add_slide(content_layout)
        sl.shapes.title.text = s["heading"]
        sl.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(r, g, b)
        sl.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

        tf = sl.placeholders[1].text_frame
        tf.clear()
        for i, bullet in enumerate(s["bullets"]):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = bullet
            p.font.size = Pt(18)
            p.level = 0

        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]

    out = OUTPUT_DIR / filename
    prs.save(str(out))
    return {"success": True, "path": str(out), "slides_count": len(slides) + 1}


def tool_create_word_report(args: dict) -> dict:
    if not DOCX_AVAILABLE:
        return {"error": "python-docx לא מותקן. הרץ: pip install python-docx"}

    filename = args["filename"].rstrip(".docx") + ".docx"
    doc = Document()

    # כותרת ראשית
    h = doc.add_heading(args["title"], level=0)
    h.runs[0].font.color.rgb = DocxRGB(0x1F, 0x4E, 0x79)

    doc.add_paragraph(f"נוצר: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph("")

    for sec in args["sections"]:
        doc.add_heading(sec["heading"], level=1)
        doc.add_paragraph(sec["content"])

        if sec.get("table"):
            tbl_data = sec["table"]
            headers  = tbl_data.get("headers", [])
            rows     = tbl_data.get("rows", [])
            if headers:
                table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                table.style = "Light Shading Accent 1"
                hdr_cells = table.rows[0].cells
                for i, h_text in enumerate(headers):
                    hdr_cells[i].text = h_text
                for ri, row in enumerate(rows):
                    for ci, cell_val in enumerate(row):
                        table.rows[ri + 1].cells[ci].text = str(cell_val)
        doc.add_paragraph("")

    out = OUTPUT_DIR / filename
    doc.save(str(out))
    return {"success": True, "path": str(out)}


def tool_create_statistics_chart(args: dict) -> dict:
    if not MATPLOTLIB_AVAILABLE:
        return {"error": "matplotlib לא מותקן. הרץ: pip install matplotlib"}

    chart_type = args["chart_type"]
    labels     = args["labels"]
    values     = args["values"]
    title      = args["title"]
    filename   = args["filename"].rstrip(".png") + ".png"
    colors     = args.get("colors") or None

    fig, ax = plt.subplots(figsize=(10, 6))
    fig.patch.set_facecolor("#F9F9F9")
    ax.set_facecolor("#F9F9F9")

    if chart_type == "bar":
        bars = ax.bar(labels, values, color=colors or "#1F4E79", edgecolor="white")
        ax.bar_label(bars, fmt="%.1f", padding=3)
        ax.set_xlabel(args.get("xlabel", ""))
        ax.set_ylabel(args.get("ylabel", ""))

    elif chart_type == "horizontal_bar":
        bars = ax.barh(labels, values, color=colors or "#1F4E79", edgecolor="white")
        ax.bar_label(bars, fmt="%.1f", padding=3)
        ax.set_xlabel(args.get("xlabel", ""))

    elif chart_type == "pie":
        ax.pie(values, labels=labels, colors=colors, autopct="%1.1f%%",
               startangle=140, wedgeprops={"edgecolor": "white"})
        ax.axis("equal")

    elif chart_type == "line":
        ax.plot(labels, values, marker="o", color=colors[0] if colors else "#1F4E79",
                linewidth=2.5, markersize=7)
        ax.fill_between(range(len(labels)), values, alpha=0.1, color="#1F4E79")
        ax.set_xlabel(args.get("xlabel", ""))
        ax.set_ylabel(args.get("ylabel", ""))

    ax.set_title(title, fontsize=14, fontweight="bold", pad=15)
    plt.tight_layout()

    out = OUTPUT_DIR / filename
    plt.savefig(str(out), dpi=150, bbox_inches="tight")
    plt.close(fig)
    return {"success": True, "path": str(out)}


def tool_recall_memory(memory: dict) -> dict:
    stats = memory["stats"]
    recent_errors = [
        {"task": e["task"][:100], "error": e["error"][:150]}
        for e in memory["errors"][-5:]
    ]
    recent_sessions = [
        {"ts": s["ts"], "task": s["task"][:80], "status": s["status"]}
        for s in memory["sessions"][-5:]
    ]
    return {
        "stats": stats,
        "recent_errors": recent_errors,
        "recent_sessions": recent_sessions,
        "improvements": memory.get("improvements", [])[-5:]
    }


# ═══════════════════════════════════════════════
#  לולאת הסוכן
# ═══════════════════════════════════════════════

def execute_tool(name: str, args: dict, memory: dict) -> str:
    try:
        if name == "create_presentation":
            result = tool_create_presentation(args)
        elif name == "create_word_report":
            result = tool_create_word_report(args)
        elif name == "create_statistics_chart":
            result = tool_create_statistics_chart(args)
        elif name == "recall_memory":
            result = tool_recall_memory(memory)
        else:
            result = {"error": f"כלי לא מוכר: {name}"}

        if "error" in result:
            log_error(memory, name, result["error"])
        return json.dumps(result, ensure_ascii=False)

    except Exception as exc:
        err = traceback.format_exc()
        log_error(memory, name, str(exc), err[:300])
        return json.dumps({"error": str(exc), "traceback": err[:300]}, ensure_ascii=False)


def run_moshe_ai(user_request: str, verbose: bool = True) -> str:
    """מריץ את MosheAI על בקשת משתמש ומחזיר תגובה טקסטואלית."""
    memory = load_memory()
    client = anthropic.Anthropic()  # ANTHROPIC_API_KEY מהסביבה

    # הוסף לקחים קודמים ל-system
    lessons = build_lessons_prompt(memory)
    system = SYSTEM_PROMPT
    if lessons:
        system += f"\n\n{lessons}"

    messages = [{"role": "user", "content": user_request}]

    if verbose:
        print(f"\n{'='*60}")
        print(f"🤖 MosheAI | בקשה: {user_request[:80]}")
        print(f"{'='*60}")

    try:
        # לולאת agent
        while True:
            with client.messages.stream(
                model=MODEL,
                max_tokens=8192,
                thinking={"type": "adaptive"},
                system=system,
                tools=TOOLS,
                messages=messages
            ) as stream:
                response = stream.get_final_message()

            if verbose:
                for block in response.content:
                    if block.type == "text" and block.text:
                        print(f"\n💬 {block.text}")

            # אם אין קריאות כלים - סיימנו
            tool_uses = [b for b in response.content if b.type == "tool_use"]
            if not tool_uses:
                break

            # הוסף תגובת assistant להיסטוריה
            messages.append({"role": "assistant", "content": response.content})

            # הרץ כלים
            tool_results = []
            for tu in tool_uses:
                if verbose:
                    print(f"\n🔧 מפעיל כלי: {tu.name}")
                result_str = execute_tool(tu.name, tu.input, memory)
                result_data = json.loads(result_str)
                if verbose:
                    if "path" in result_data:
                        print(f"   ✅ נשמר: {result_data['path']}")
                    elif "error" in result_data:
                        print(f"   ❌ שגיאה: {result_data['error']}")
                tool_results.append({
                    "type": "tool_result",
                    "tool_use_id": tu.id,
                    "content": result_str
                })

            messages.append({"role": "user", "content": tool_results})

        # תגובה סופית
        final_text = next(
            (b.text for b in response.content if b.type == "text" and b.text),
            "המשימה הושלמה."
        )

        # בדוק אם יש פלטי קבצים
        outputs = []
        for msg in messages:
            if isinstance(msg.get("content"), list):
                for block in msg["content"]:
                    if isinstance(block, dict) and block.get("type") == "tool_result":
                        try:
                            d = json.loads(block["content"])
                            if d.get("path"):
                                outputs.append(d["path"])
                        except Exception:
                            pass

        log_success(memory, user_request, "; ".join(outputs))
        return final_text

    except anthropic.AuthenticationError:
        err = "שגיאת API Key - ודא שמשתנה הסביבה ANTHROPIC_API_KEY מוגדר."
        log_error(memory, user_request, err)
        return err
    except Exception as exc:
        err = traceback.format_exc()
        log_error(memory, user_request, str(exc), err[:300])
        return f"שגיאה: {exc}"


# ═══════════════════════════════════════════════
#  CLI אינטראקטיבי
# ═══════════════════════════════════════════════

HELP_TEXT = """
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
🤖  MosheAI - עוזר חכם לדוחות ושקופיות
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
פקודות מיוחדות:
  /memory   - הצג זיכרון והיסטוריה
  /stats    - הצג סטטיסטיקת שימוש
  /outputs  - פתח תיקיית הפלטים
  /help     - הצג עזרה
  /quit     - יציאה

דוגמאות לבקשות:
  • צור מצגת בנושא ביצועי מכירות לרבעון האחרון
  • כתוב דוח Word על מצב מלאי המחסן
  • צור גרף עמודות של הכנסות לפי חודש עם הנתונים: ינואר=100, פברואר=150, מרץ=130
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
"""


def main():
    print(HELP_TEXT)
    memory = load_memory()
    s = memory["stats"]
    print(f"📊 סטטיסטיקה: {s['total_tasks']} משימות | ✅ {s['success']} הצליחו | ❌ {s['failed']} נכשלו")
    print(f"📁 פלטים בתיקייה: {OUTPUT_DIR}\n")

    while True:
        try:
            user_input = input("👤 הזן בקשה: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nשלום!")
            break

        if not user_input:
            continue

        if user_input == "/quit":
            print("שלום!")
            break
        elif user_input == "/help":
            print(HELP_TEXT)
        elif user_input == "/memory":
            m = load_memory()
            print(json.dumps(tool_recall_memory(m), ensure_ascii=False, indent=2))
        elif user_input == "/stats":
            s = load_memory()["stats"]
            print(f"סה\"כ: {s['total_tasks']} | הצלחות: {s['success']} | כישלונות: {s['failed']}")
        elif user_input == "/outputs":
            os.startfile(str(OUTPUT_DIR))
        else:
            run_moshe_ai(user_input, verbose=True)


if __name__ == "__main__":
    main()
